import codecs
import csv
import hashlib
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any
import pandas as pd
from openpyxl import load_workbook
from fastapi import HTTPException
from psycopg import connection
from pypdf import PdfReader

from app.worker_manager import ensure_csv_worker

from ..config import FILE_STORAGE_DIR, TESSDATA_DIR, TESSERACT_CMD, TITLE_HEADER_WORDS, CONTENT_HEADER_WORDS
from ..ai import fallback

def row_to_complaint(row: dict[str, Any], source_row: int) -> dict[str, Any] | None:
    headers = list(row)
    if "해결명(SOLUTION_CRTR_NAME)" in headers and "분쟁유형명(DISPUTE_TYPE_NAME)" in headers:
        item, dispute, solution = str(row.get("품목명(ITEM_NAME)", "")).strip(), str(row.get("분쟁유형명(DISPUTE_TYPE_NAME)", "")).strip(), str(row.get("해결명(SOLUTION_CRTR_NAME)", "")).strip()
        return {"source_row": source_row, **fallback(" · ".join(filter(None, [item, dispute])), f"{dispute}\n{solution}")} if solution else None
    def column(words: list[str]) -> str | None:
        return next((header for header in headers if any(re.sub(r"[\s_·-]+", "", word).lower() in re.sub(r"[\s_·-]+", "", str(header)).lower() for word in words)), None)
    title_key, content_key = column(["민원 제목", "제목", "title", "subject"]), column(["민원 내용", "민원내용", "신청원인", "원문", "내용", "complaint", "content", "질문"])
    content = str(row.get(content_key, "")).strip() if content_key else ""
    return {"source_row": source_row, **fallback(str(row.get(title_key, "")) if title_key else "", content)} if content else None

def csv_schema_signature(headers: list[str]) -> str:
    normalized = "\x1f".join(re.sub(r"\s+", "", header).lower() for header in headers)
    return hashlib.sha256(normalized.encode()).hexdigest()

def csv_preview(path: Path, encoding: str, size: int = 5) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.DictReader(source)
        headers = [header.strip() for header in (reader.fieldnames or []) if header and header.strip()]
        samples = [{key: str(value or "").strip()[:500] for key, value in row.items() if key} for _, row in zip(range(size), reader)]
    return headers, samples

def default_csv_mapping(headers: list[str]) -> dict[str, Any]:
    def column(words: list[str]) -> str:
        return next((header for header in headers if any(re.sub(r"[\s_·-]+", "", word).lower() in re.sub(r"[\s_·-]+", "", header).lower() for word in words)), "")
    if "해결명(SOLUTION_CRTR_NAME)" in headers and "분쟁유형명(DISPUTE_TYPE_NAME)" in headers:
        return {"title_column": "품목명(ITEM_NAME)", "content_columns": ["분쟁유형명(DISPUTE_TYPE_NAME)", "해결명(SOLUTION_CRTR_NAME)"], "response_column": "", "category_column": "", "confidence": 0.95, "reason": "공정거래위원회 상담 사례 헤더를 인식했습니다.", "source": "profile"}
    title_column, content_column = column(TITLE_HEADER_WORDS), column(CONTENT_HEADER_WORDS)
    return {"title_column": title_column, "content_columns": [content_column] if content_column else [], "response_column": "", "category_column": "", "confidence": 0.85 if content_column else 0.0, "reason": "기본 민원 헤더 후보를 인식했습니다." if content_column else "자동으로 민원 본문 열을 찾지 못했습니다.", "source": "heuristic"}

def validate_csv_mapping(mapping: dict[str, Any], headers: list[str], samples: list[dict[str, Any]]) -> dict[str, Any]:
    allowed_headers = set(headers)
    title_column = str(mapping.get("title_column") or "").strip()
    content_columns = [str(value).strip() for value in mapping.get("content_columns", []) if str(value).strip() in allowed_headers]
    response_column = str(mapping.get("response_column") or "").strip()
    category_column = str(mapping.get("category_column") or "").strip()
    if title_column not in allowed_headers:
        title_column = ""
    if response_column not in allowed_headers:
        response_column = ""
    if category_column not in allowed_headers:
        category_column = ""
    content_columns = list(dict.fromkeys(content_columns))
    nonempty = [" ".join(str(row.get(column, "")).strip() for column in content_columns).strip() for row in samples]
    nonempty_ratio = sum(bool(value) for value in nonempty) / max(len(samples), 1)
    average_length = sum(len(value) for value in nonempty) / max(len(nonempty), 1)
    title_average = sum(len(str(row.get(title_column, "")).strip()) for row in samples) / max(len(samples), 1) if title_column else 0
    supplied_confidence = max(0.0, min(1.0, float(mapping.get("confidence", 0))))
    validation_score = min(1.0, nonempty_ratio * 0.55 + min(average_length / 160, 1.0) * 0.35 + (0.10 if not title_column or title_average <= 180 else 0.0))
    return {"title_column": title_column, "content_columns": content_columns, "response_column": response_column, "category_column": category_column, "confidence": round(min(supplied_confidence, validation_score) if supplied_confidence else validation_score, 2), "reason": str(mapping.get("reason") or ""), "source": str(mapping.get("source") or "manual"), "valid": bool(content_columns and nonempty_ratio >= 0.6 and average_length >= 8)}

def csv_row_to_record(row: dict[str, Any], source_row: int, mapping: dict[str, Any]) -> dict[str, Any] | None:
    content_parts = [str(row.get(column, "")).strip() for column in mapping["content_columns"]]
    content = "\n".join(part for part in content_parts if part)
    if not content:
        return None
    title = str(row.get(mapping.get("title_column", ""), "")).strip()
    return {"source_row": source_row, "title": title, "content": content, "source_response": str(row.get(mapping.get("response_column", ""), "")).strip() if mapping.get("response_column") else "", "source_category": str(row.get(mapping.get("category_column", ""), "")).strip() if mapping.get("category_column") else ""}

def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()

def _is_complaint_header(values: list[Any]) -> bool:
    headers = [re.sub(r"\s+", "", _cell_text(value)).lower() for value in values]
    has_title = any(value in {"민원제목", "제목", "title", "subject"} for value in headers)
    has_content = any(value in {"민원내용", "신청원인", "원문", "내용", "complaint", "content", "질문"} for value in headers)
    return has_title and has_content

def _records_from_rows(rows: list[tuple[int, list[Any]]], sheet_name: str = "") -> list[dict[str, Any]]:
    header_at = next((index for index, (_, values) in enumerate(rows[:30]) if _is_complaint_header(values)), None)
    if header_at is None:
        return []
    _, header_values = rows[header_at]
    headers = [_cell_text(value) for value in header_values]
    records: list[dict[str, Any]] = []
    for source_row, values in rows[header_at + 1:]:
        row = {headers[index]: _cell_text(value) for index, value in enumerate(values) if index < len(headers) and headers[index]}
        record = row_to_complaint(row, source_row)
        if record:
            record["source_sheet"] = sheet_name
            records.append(record)
    return records

def spreadsheet_records(path: Path, suffix: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                rows = [(number, list(values)) for number, values in enumerate(worksheet.iter_rows(values_only=True), start=1)]
                records.extend(_records_from_rows(rows, worksheet.title))
        finally:
            workbook.close()
    else:
        sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
        for sheet_name, frame in sheets.items():
            rows = [(number, row) for number, row in enumerate(frame.fillna("").values.tolist(), start=1)]
            records.extend(_records_from_rows(rows, str(sheet_name)))
    return records

def _text_quality(text: str, required_markers: tuple[str, ...] = ()) -> float:
    compact = re.sub(r"\s+", "", text or "")
    if not compact:
        return 0.0
    readable = sum(character.isalnum() or "가" <= character <= "힣" or character in ".,!?·()[]'\"-:/" for character in compact)
    score = min(len(compact) / 160, 1.0) * 0.45 + (readable / len(compact)) * 0.35
    score += 0.20 * (sum(marker in text for marker in required_markers) / max(len(required_markers), 1))
    return round(min(score, 1.0), 3)

def _installed_tool(name: str) -> str | None:
    found = shutil.which(name)
    if found:
        return found
    extension = ".exe" if sys.platform == "win32" else ""
    candidates = [Path(sys.executable).with_name(f"{name}{extension}")]
    if name == "soffice" and sys.platform == "win32":
        candidates.extend([
            Path("C:/Program Files/LibreOffice/program/soffice.exe"),
            Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
        ])
    if name == "tesseract" and sys.platform == "win32":
        candidates.extend([
            Path("C:/Program Files/Tesseract-OCR/tesseract.exe"),
            Path("C:/Program Files (x86)/Tesseract-OCR/tesseract.exe"),
        ])
    if name == "hwp5txt" and sys.platform == "win32":
        candidates.append(Path(__file__).resolve().parents[2] / ".hwp-parser" / "Scripts" / "hwp5txt.exe")
    return next((str(candidate) for candidate in candidates if candidate.is_file()), None)

def ocr_image(image: Any, pytesseract_module: Any | None = None) -> str:
    if not TESSDATA_DIR.is_dir() or not (TESSDATA_DIR / "kor.traineddata").is_file() or not (TESSDATA_DIR / "eng.traineddata").is_file():
        raise RuntimeError("Tesseract 한국어·영어 언어 데이터가 준비되지 않았습니다.")
    if pytesseract_module is None:
        import pytesseract as pytesseract_module
    executable = TESSERACT_CMD or _installed_tool("tesseract")
    if not executable:
        raise RuntimeError("Tesseract 실행 파일을 찾지 못했습니다.")
    pytesseract_module.pytesseract.tesseract_cmd = executable
    return pytesseract_module.image_to_string(image, lang="kor+eng", config=f"--tessdata-dir {TESSDATA_DIR}")

def _ocr_pdf_page(path: Path, page_index: int) -> str:
    try:
        import fitz
        from PIL import Image
        import pytesseract
        document = fitz.open(path)
        page = document.load_page(page_index)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
        return ocr_image(image, pytesseract)
    except Exception:
        return ""

def _pdf_page_candidates(path: Path) -> list[tuple[str, float, float]]:
    reader = PdfReader(path)
    native = [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]
    plumber_text = [""] * len(native)
    image_coverage = [0.0] * len(native)
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for index, page in enumerate(pdf.pages):
                plumber_text[index] = page.extract_text() or ""
                page_area = max(float(page.width * page.height), 1.0)
                image_coverage[index] = min(1.0, sum(float(item.get("width", 0) * item.get("height", 0)) for item in page.images) / page_area)
    except Exception:
        pass

    selected: list[tuple[str, float, float]] = []
    for index, primary in enumerate(native):
        alternate = plumber_text[index]
        primary_score = _text_quality(primary, ("신청원인",))
        alternate_score = _text_quality(alternate, ("신청원인",))
        text, score = (alternate, alternate_score) if alternate_score > primary_score else (primary, primary_score)
        if score < 0.62 or (image_coverage[index] >= 0.70 and score < 0.82):
            ocr_text = _ocr_pdf_page(path, index)
            ocr_score = _text_quality(ocr_text, ("신청원인",))
            if ocr_score > score:
                text, score = ocr_text, ocr_score
        selected.append((text, score, image_coverage[index]))
    return selected

def _records_from_case_pages(pages: list[tuple[str, float, float]], filename: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for page_number, (text, _score, _coverage) in enumerate(pages, start=1):
        heading = re.search(r"사례\s*(\d{3})", text)
        marker = re.search(r"신청원인\s*", text)
        if not heading or not marker:
            continue
        before = text[:marker.start()].splitlines()
        title_lines = [line.strip() for line in before if line.strip() and not re.search(r"사례\s*\d{3}", line)]
        body = text[marker.end():]
        body = re.split(r"피신청인\s*등의\s*주장|가상\s*민원\s*데이터", body, maxsplit=1)[0].strip()
        title = " ".join(title_lines).strip()
        if title and body:
            records.append({"source_row": page_number, "source_case": heading.group(1), **fallback(title, body)})
    if records:
        return records
    combined = "\n".join(text for text, _, _ in pages).strip()
    return [{"source_row": 1, **fallback(Path(filename).stem, combined)}] if combined else []

def _records_from_hwp_text(text: str, filename: str) -> list[dict[str, Any]]:
    headings = list(re.finditer(r"(?m)^\s*\d{2}\s+.+?/\s*사례\s*(\d{3})\s*$", text))
    records: list[dict[str, Any]] = []
    for index, heading in enumerate(headings):
        block = text[heading.end(): headings[index + 1].start() if index + 1 < len(headings) else len(text)]
        marker = re.search(r"신청원인\s*", block)
        if not marker:
            continue
        title = " ".join(line.strip() for line in block[:marker.start()].splitlines() if line.strip())
        body = re.split(r"피신청인\s*등의\s*주장|가상\s*민원\s*데이터", block[marker.end():], maxsplit=1)[0].strip()
        if title and body:
            records.append({"source_row": int(heading.group(1)), "source_case": heading.group(1), **fallback(title, body)})
    if records:
        return records
    return _records_from_case_pages([(text, _text_quality(text, ("신청원인",)), 0.0)], filename)

def pdf_records(path: Path, filename: str) -> list[dict[str, Any]]:
    return _records_from_case_pages(_pdf_page_candidates(path), filename)

def _hwp_to_pdf_records(path: Path, filename: str) -> list[dict[str, Any]]:
    converter = _installed_tool("soffice")
    if not converter:
        return []
    conversion_root = FILE_STORAGE_DIR / "conversion-temp"
    conversion_root.mkdir(parents=True, exist_ok=True)
    try:
        with tempfile.TemporaryDirectory(prefix="complaintai-hwp-", dir=conversion_root) as output_dir:
            converted = subprocess.run([converter, "--headless", "--convert-to", "pdf", "--outdir", output_dir, str(path)], capture_output=True, check=False)
            pdf_path = Path(output_dir) / f"{path.stem}.pdf"
            return pdf_records(pdf_path, filename) if not converted.returncode and pdf_path.exists() else []
    except OSError:
        return []

def hwp_records(path: Path, filename: str) -> list[dict[str, Any]]:
    extractor = _installed_tool("hwp5txt")
    text = ""
    if extractor:
        completed = subprocess.run([extractor, str(path)], capture_output=True, check=False)
        if not completed.returncode:
            text = completed.stdout.decode("utf-8", errors="replace").strip()
    confidence = _text_quality(text, ("신청원인",))
    if text and confidence >= 0.62:
        return _records_from_hwp_text(text, filename)
    converted_records = _hwp_to_pdf_records(path, filename)
    if converted_records:
        return converted_records
    if text:
        return _records_from_hwp_text(text, filename)
    raise HTTPException(503, "HWP 텍스트 추출기(hwp5txt/pyhwp) 또는 HWP-to-PDF 변환기가 설치되어 있지 않습니다.")

def csv_encoding(path: Path) -> str:
    # Read a bounded sample; do not load a potentially 1GB CSV into memory.
    with path.open("rb") as source:
        sample = source.read(65536)
    for marker, encoding in ((codecs.BOM_UTF8, "utf-8-sig"),
                             (codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
                             (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16")):
        if sample.startswith(marker):
            return encoding
    try:
        # An incomplete character at the sample boundary is not invalid UTF-8.
        codecs.getincrementaldecoder("utf-8")().decode(sample, final=False)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp949"

def csv_rows(path: Path, encoding: str):
    with path.open("r", encoding=encoding, newline="") as source:
        yield from csv.DictReader(source)

def start_import_worker(job_id: str) -> None:
    try:
        ensure_csv_worker()
    except Exception as error:
        message = "CSV 처리 워커를 실행하지 못했습니다. 재처리를 눌러 다시 시도하거나 서버 실행 환경을 확인해 주세요."
        with connection() as db:
            with db.cursor() as cur:
                cur.execute("UPDATE import_jobs SET status='failed',last_error=%s WHERE id=%s AND status='queued'", (message, job_id))
            db.commit()
        raise HTTPException(503, message) from error